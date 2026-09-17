"""DOCX import/export via python-docx (lossy but practical bridge)."""
from __future__ import annotations

import io
import re
import uuid

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor

from .models import Block, BlockType, Document, SectionEntry, TableData


def _id() -> str:
    return uuid.uuid4().hex[:12]


def _clean_runs(paragraph) -> str:
    """Rebuild inline **bold** / *italic* markers from docx runs."""
    parts = []
    for run in paragraph.runs:
        t = run.text
        if not t:
            continue
        if run.bold and run.italic:
            parts.append(f"***{t}***")
        elif run.bold:
            parts.append(f"**{t}**")
        elif run.italic:
            parts.append(f"*{t}*")
        else:
            parts.append(t)
    return "".join(parts)


def document_to_docx(doc: Document) -> bytes:
    d = DocxDocument()
    core = d.core_properties
    core.title = doc.title
    if doc.company_name:
        core.author = doc.company_name
    d.add_heading(doc.title, level=0)
    for b in doc.blocks:
        if b.type == BlockType.heading:
            d.add_heading(re.sub(r"[*`]", "", b.text), level=min(3, b.level))
        elif b.type in (BlockType.paragraph, BlockType.quote):
            d.add_paragraph(re.sub(r"[*`]", "", b.text))
        elif b.type == BlockType.code:
            p = d.add_paragraph()
            run = p.add_run(b.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif b.type == BlockType.divider:
            d.add_paragraph("─" * 40)
        elif b.type == BlockType.keyvalue:
            for r in b.rows:
                p = d.add_paragraph()
                run = p.add_run(f"{r.label}: " if r.label else "")
                run.bold = True
                p.add_run(r.value)
        elif b.type in (BlockType.bullets, BlockType.checklist):
            items = b.items if b.type == BlockType.bullets else [
                f"{'[x]' if i.checked else '[ ]'} {i.text}" for i in b.checklist]
            for item in items:
                if item.strip():
                    d.add_paragraph(re.sub(r"[*`]", "", item), style="List Bullet")
        elif b.type == BlockType.sections:
            for n, e in enumerate(b.entries, 1):
                d.add_paragraph(f"{n}. {re.sub(r'[*`]', '', e.title)}", style="List Number")
                if e.text.strip():
                    d.add_paragraph(re.sub(r"[*`]", "", e.text))
        elif b.type == BlockType.table:
            headers = [h for h in b.table.headers]
            rows = [r for r in b.table.rows if any(c.strip() for c in r)]
            ncols = max([len(headers)] + [len(r) for r in rows] + [0])
            if ncols:
                table = d.add_table(rows=1 + len(rows), cols=ncols)
                table.style = "Table Grid"
                for j in range(ncols):
                    cell = table.cell(0, j)
                    cell.text = headers[j] if j < len(headers) else ""
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.bold = True
                for i, r in enumerate(rows, 1):
                    for j in range(ncols):
                        table.cell(i, j).text = r[j] if j < len(r) else ""
        elif b.type == BlockType.signatures:
            p = d.add_paragraph()
            p.add_run(f"{b.left.name} ({b.left.role})".strip())
            p.add_run("\t\t")
            p.add_run(f"{b.right.name} ({b.right.role})".strip())
        elif b.type == BlockType.columns:
            cols = [(c.title, c.text) for c in b.columns if c.title.strip() or c.text.strip()]
            if cols:
                table = d.add_table(rows=2, cols=len(cols))
                table.style = "Table Grid"
                for j, (title, text) in enumerate(cols):
                    cell = table.cell(0, j)
                    cell.text = title
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.bold = True
                    table.cell(1, j).text = text
        elif b.type == BlockType.pagebreak:
            d.add_page_break()
        elif b.type == BlockType.toc:
            d.add_paragraph("Table of contents")
            for e in b.toc_entries:
                d.add_paragraph(f"{e.title} … {e.page}")
        elif b.type == BlockType.image:
            if b.image.caption.strip():
                p = d.add_paragraph()
                run = p.add_run(b.image.caption)
                run.italic = True
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _has_page_break(paragraph) -> bool:
    """Detect <w:br w:type="page"/> inside a paragraph's runs."""
    for run in paragraph.runs:
        for br in run._element.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br"):
            if br.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type") == "page":
                return True
    return False


def docx_to_document(data: bytes, title: str = "Imported document") -> Document:
    d = DocxDocument(io.BytesIO(data))
    blocks: list[Block] = []
    doc_title = title
    first = True
    for p in d.paragraphs:
        if _has_page_break(p):
            blocks.append(Block(id=_id(), type=BlockType.pagebreak))
            first = False
            continue
        text = _clean_runs(p).strip()
        if not text:
            continue
        style = (p.style.name or "").lower()
        if first and style.startswith("title"):
            doc_title = re.sub(r"[*`]", "", text) or title
            first = False
            continue
        first = False
        if style.startswith("heading"):
            try:
                level = int(style.replace("heading", "").strip() or "1")
            except ValueError:
                level = 1
            blocks.append(Block(id=_id(), type=BlockType.heading, level=max(1, min(3, level)), text=text))
        elif "bullet" in style or "list bullet" in style:
            if blocks and blocks[-1].type == BlockType.bullets:
                blocks[-1].items.append(text)
            else:
                blocks.append(Block(id=_id(), type=BlockType.bullets, items=[text]))
        elif "number" in style:
            m = re.match(r"^(\d+)\.\s+(.*)$", text)
            entry = SectionEntry(title=(m.group(2) if m else text), text="")
            if blocks and blocks[-1].type == BlockType.sections:
                blocks[-1].entries.append(entry)
            else:
                blocks.append(Block(id=_id(), type=BlockType.sections, entries=[entry]))
        else:
            blocks.append(Block(id=_id(), type=BlockType.paragraph, text=text))
    for table in d.tables:
        grid = [[c.text.strip() for c in row.cells] for row in table.rows]
        grid = [r for r in grid if any(r)]
        if grid:
            blocks.append(Block(id=_id(), type=BlockType.table,
                                table=TableData(headers=grid[0], rows=grid[1:])))
    return Document(title=doc_title, blocks=blocks)


__all__ = ["document_to_docx", "docx_to_document"]
